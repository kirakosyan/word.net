"""Author synthetic DrawingML presets independently with python-docx."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import docx
from docx.oxml import parse_xml
from docx.shared import Inches

PRESETS = ('rect', 'triangle', 'diamond')


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def create(path, geometry):
    document = docx.Document()
    document.core_properties.author = 'Synthetic fixture'
    document.core_properties.created = document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(0.5)
    document.add_paragraph('Before')
    drawing = parse_xml('''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
 <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
  <wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>1270000</wp:posOffset></wp:positionH>
  <wp:positionV relativeFrom="page"><wp:posOffset>1270000</wp:posOffset></wp:positionV><wp:extent cx="1778000" cy="1016000"/>
  <wp:wrapNone/><wp:docPr id="1" name="Synthetic preset"/><a:graphic>
   <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr/><wps:spPr>
    <a:xfrm><a:off x="0" y="0"/><a:ext cx="1778000" cy="1016000"/></a:xfrm>
    <a:prstGeom prst="%s"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="2468AC"/></a:solidFill><a:ln><a:noFill/></a:ln>
   </wps:spPr><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic>
 </wp:anchor></w:drawing>''' % geometry)
    document.add_paragraph().add_run()._r.append(drawing)
    document.add_paragraph('After')
    temporary = path.with_suffix('.unpacked.docx')
    document.save(temporary)
    with ZipFile(temporary) as source, ZipFile(path, 'w') as output:
        for name in source.namelist():
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            output.writestr(info, source.read(name), compress_type=ZIP_DEFLATED)
    temporary.unlink()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    directory = parser.parse_args().output.resolve()
    fixture_root = (Path(__file__).resolve().parents[1] / 'tests/wordnet.Conversion.Tests/Fixtures').resolve()
    if directory == fixture_root or fixture_root in directory.parents or directory.exists():
        parser.error('Choose a new scratch directory outside the fixture tree.')
    if docx.__version__ != '1.2.0':
        raise RuntimeError('This fixture requires python-docx 1.2.0.')
    directory.mkdir(parents=True)
    result = {}
    for geometry in PRESETS:
        path = directory / (geometry + '.docx')
        create(path, geometry)
        result[path.name] = digest(path)
    print(json.dumps({'producer': 'python-docx 1.2.0 (MIT), synthetic DrawingML preset shapes',
                      'generatorSha256': digest(Path(__file__)), 'documents': result}))


if __name__ == '__main__':
    main()
