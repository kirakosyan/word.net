"""Author synthetic vector anchors and explicit alternatives with python-docx."""
import argparse
import copy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches

DESCRIPTION = 'Blue trend diagram'
KINDS = ('line', 'ellipse', 'roundRect', 'rect', 'triangle', 'diamond')


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def create(path, mode):
    document = docx.Document()
    document.core_properties.author = 'Synthetic fixture'
    document.core_properties.title = 'Synthetic accessible shapes'
    document.core_properties.created = document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    section = document.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(0.5)
    before = document.add_paragraph('Before')
    if mode == 'selection':
        br = OxmlElement('w:br'); br.set(qn('w:type'), 'page'); before.add_run()._r.append(br)
    prefix = mode.rsplit('-', 1)[0] if '-' in mode else ''
    geometry = prefix if prefix in KINDS else 'triangle'
    decorative = mode == 'decorative' or mode.endswith('-decorative')
    description = '' if decorative else ' descr="' + DESCRIPTION + '"'
    stroke = '<a:solidFill><a:srgbClr val="FF0000"/></a:solidFill>' if mode == 'stroke' or geometry == 'line' else '<a:noFill/>'
    fill = '<a:noFill/>' if geometry == 'line' else '<a:solidFill><a:srgbClr val="2468AC"/></a:solidFill>'
    drawing = parse_xml('''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
 <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
  <wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>1270000</wp:posOffset></wp:positionH>
  <wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV><wp:extent cx="1778000" cy="1016000"/>
  <wp:wrapTopAndBottom/><wp:docPr id="1" name="Synthetic preset"%s/><a:graphic>
   <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr/><wps:spPr>
    <a:xfrm><a:off x="0" y="0"/><a:ext cx="1778000" cy="1016000"/></a:xfrm>
    <a:prstGeom prst="%s"><a:avLst/></a:prstGeom>%s<a:ln w="25400">%s</a:ln>
   </wps:spPr><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic>
 </wp:anchor></w:drawing>''' % (description, geometry, fill, stroke))
    document.add_paragraph().add_run()._r.append(drawing)
    if mode == 'two':
        second = copy.deepcopy(drawing)
        properties = second.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr')
        properties.set('id', '2'); properties.set('descr', 'Second blue diagram')
        document.add_paragraph().add_run()._r.append(second)
    if mode in ('box', 'box-decorative'):
        box = parse_xml('''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
 <wp:anchor simplePos="0" relativeHeight="0" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
  <wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="margin"><wp:posOffset>0</wp:posOffset></wp:positionH>
  <wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV><wp:extent cx="2540000" cy="1016000"/>
  <wp:wrapTopAndBottom/><wp:docPr id="2" name="Synthetic box"/><a:graphic>
   <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></wps:spPr>
    <wps:txbx><w:txbxContent><w:p><w:r><w:t>Box text</w:t></w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr><a:noAutofit/></wps:bodyPr>
   </wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>''')
        document.add_paragraph().add_run()._r.append(box)
    document.add_paragraph('After')
    temporary = path.with_suffix('.unpacked.docx'); document.save(temporary)
    with ZipFile(temporary) as source, ZipFile(path, 'w') as output:
        for name in source.namelist():
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            output.writestr(info, source.read(name), compress_type=ZIP_DEFLATED)
    temporary.unlink()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    directory = parser.parse_args().output.resolve()
    fixtures = (Path(__file__).resolve().parents[1] / 'tests/wordnet.Conversion.Tests/Fixtures').resolve()
    if directory == fixtures or fixtures in directory.parents or directory.exists():
        parser.error('Choose a new scratch directory outside the fixture tree.')
    if docx.__version__ != '1.2.0':
        raise RuntimeError('This fixture requires python-docx 1.2.0.')
    directory.mkdir(parents=True)
    documents = {}
    modes = ['described', 'decorative', 'selection', 'two', 'stroke', 'box', 'box-decorative']
    modes.extend(kind + '-' + policy for kind in KINDS for policy in ('described', 'decorative'))
    for mode in modes:
        path = directory / (mode + '.docx'); create(path, mode); documents[path.name] = digest(path)
    print(json.dumps({'producer': 'python-docx 1.2.0 (MIT), synthetic described and decorative vectors',
                      'generatorSha256': digest(Path(__file__)), 'documents': documents}))


if __name__ == '__main__':
    main()
