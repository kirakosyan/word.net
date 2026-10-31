"""Create a synthetic narrow-page DOCX and a locally verified English break entry."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pyphen


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    directory = parser.parse_args().output.resolve()
    fixtures = (Path(__file__).resolve().parents[1] / 'tests/wordnet.Conversion.Tests/Fixtures').resolve()
    if directory.exists() or directory == fixtures or fixtures in directory.parents:
        parser.error('Choose a new scratch directory outside the fixture tree.')
    if docx.__version__ != '1.2.0' or pyphen.__version__ != '0.17.2':
        raise RuntimeError('Requires python-docx 1.2.0 and Pyphen 0.17.2.')
    if pyphen.Pyphen(lang='en_US').inserted('extraordinary') != 'ex-tra-or-di-nary':
        raise ValueError('The independent US English dictionary points changed.')
    directory.mkdir(parents=True)
    for name, width, value, justified in [
        ('source.docx', 120, 'A extraordinary example', False),
        ('justified-reference.docx', 150, 'A red ex\u00adtra\u00ador\u00addi\u00adnary example', True),
    ]:
        document = docx.Document(); document.core_properties.author = 'Synthetic fixture'
        document.core_properties.title = 'Synthetic local hyphenation'
        document.core_properties.created = document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
        section = document.sections[0]; section.page_width = Pt(width); section.page_height = Pt(792)
        section.left_margin = section.right_margin = Pt(30); section.top_margin = section.bottom_margin = Pt(36)
        paragraph = document.add_paragraph()
        if justified:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = paragraph.add_run(value); run.font.name = 'Noto Sans'; run.font.size = Pt(11)
        temporary = directory / 'raw.docx'; path = directory / name; document.save(temporary)
        with ZipFile(temporary) as source, ZipFile(path, 'w') as output:
            for member in source.namelist():
                info = ZipInfo(member, (2000, 1, 1, 0, 0, 0))
                output.writestr(info, source.read(member), compress_type=ZIP_DEFLATED)
        temporary.unlink()
    dictionary = Path(pyphen.LANGUAGES['en_US'])
    print(json.dumps({'producer': 'python-docx 1.2.0; Pyphen 0.17.2 US English pattern reference',
                      'generatorSha256': digest(Path(__file__)), 'sourceSha256': digest(directory / 'source.docx'),
                      'justifiedReferenceSha256': digest(directory / 'justified-reference.docx'),
                      'referenceDictionarySha256': digest(dictionary), 'entry': 'ex-tra-or-di-nary'}))


if __name__ == '__main__':
    main()
