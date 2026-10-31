"""Generate synthetic DOCX templates for the independent template-inspection probe.

Requires python-docx 1.2.0. No production data or Office installation is used.
"""
import argparse
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def field(parent, instruction):
    simple = OxmlElement('w:fldSimple')
    simple.set(qn('w:instr'), instruction)
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'cached'
    run.append(text); simple.append(run); parent.append(simple)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    args = parser.parse_args(); args.output.mkdir(parents=True, exist_ok=False)
    for name in ('valid', 'missing', 'type', 'repeat'):
        document = Document()
        if name == 'valid':
            field(document.add_paragraph()._p, 'MERGEFIELD Customer.Name')
            field(document.add_paragraph()._p, 'MERGEFIELD Total \\# "N2"')
        elif name == 'missing':
            field(document.add_paragraph()._p, 'MERGEFIELD HiddenAccount')
        elif name == 'type':
            field(document.add_paragraph()._p, 'MERGEFIELD Total \\# "N2"')
        else:
            control = OxmlElement('w:sdt'); properties = OxmlElement('w:sdtPr')
            tag = OxmlElement('w:tag'); tag.set(qn('w:val'), 'repeat:Items'); properties.append(tag)
            content = OxmlElement('w:sdtContent'); paragraph = OxmlElement('w:p')
            field(paragraph, 'MERGEFIELD Name'); content.append(paragraph)
            control.append(properties); control.append(content); document._body._element.insert(0, control)
        raw = args.output / (name + '.raw.docx'); result = args.output / (name + '.docx')
        document.save(raw)
        with ZipFile(raw) as source, ZipFile(result, 'w', compression=ZIP_DEFLATED) as target:
            for info in source.infolist():
                info.date_time = (2020, 1, 1, 0, 0, 0); target.writestr(info, source.read(info.filename))
        raw.unlink()


if __name__ == '__main__':
    main()
