"""Create a synthetic Hindi/document-math DOCX independently with python-docx."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--arabic', action='store_true', help='Use Arabic instead of Hindi around the equation.')
    args = parser.parse_args()
    path = args.output.resolve()
    fixture_root = (Path(__file__).resolve().parents[1] / 'tests/wordnet.Conversion.Tests/Fixtures').resolve()
    if path == fixture_root or fixture_root in path.parents or path.exists():
        parser.error('Choose a new scratch output path outside the fixture tree.')
    if docx.__version__ != '1.2.0':
        raise RuntimeError('This fixture requires python-docx 1.2.0.')
    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.core_properties.author = 'Synthetic fixture'
    document.core_properties.created = document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    normal = document.styles['Normal']
    normal.font.name = 'Noto Sans Arabic' if args.arabic else 'Noto Sans Devanagari'
    normal.font.size = Pt(11)
    word = 'مرحبا' if args.arabic else 'हिन्दी'
    document.add_paragraph('Before ' + word)
    paragraph = document.add_paragraph()
    math = OxmlElement('m:oMath')
    fraction = OxmlElement('m:f')
    for element, value in [('m:num', 'x'), ('m:den', '2')]:
        part = OxmlElement(element)
        run = OxmlElement('m:r')
        properties = OxmlElement('m:rPr')
        style = OxmlElement('m:sty')
        style.set(qn('m:val'), 'i')
        properties.append(style)
        run.append(properties)
        word_properties = OxmlElement('w:rPr')
        size = OxmlElement('w:sz')
        size.set(qn('w:val'), '24')
        word_properties.append(size)
        run.append(word_properties)
        text = OxmlElement('m:t')
        text.text = value
        run.append(text)
        part.append(run)
        fraction.append(part)
    math.append(fraction)
    paragraph._p.append(math)
    document.add_paragraph('After ' + word)
    temporary = path.with_suffix('.unpacked.docx')
    document.save(temporary)
    with ZipFile(temporary) as source, ZipFile(path, 'w') as output:
        for name in source.namelist():
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            output.writestr(info, source.read(name), compress_type=ZIP_DEFLATED)
    temporary.unlink()
    print(json.dumps({'producer': 'python-docx 1.2.0 (MIT), synthetic OMML fraction',
                      'generatorSha256': digest(Path(__file__)), 'documentSha256': digest(path)}))


if __name__ == '__main__':
    main()
