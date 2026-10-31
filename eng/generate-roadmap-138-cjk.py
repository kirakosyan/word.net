"""Author synthetic East Asian DOCX references independently with python-docx."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import docx
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

CASES = {
    'japanese-plain': ('JP', '日本語の文章です日本語の文章です'),
    'japanese': ('JP', '日本語の文章です。段落（句読点）を正しく折り返す、次の行は継続します。'),
    'simplified': ('SC', '这是一个中文段落，包含标点（以及）数字123.45。测试换行规则！'),
    'traditional': ('TC', '繁體中文測試，標點與文字換行規則。'),
}


def create(path, suffix, text):
    document = docx.Document()
    document.core_properties.author = 'Synthetic fixture'
    document.core_properties.created = document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    section = document.sections[0]; section.page_width = Pt(150); section.page_height = Inches(11)
    section.left_margin = section.right_margin = Pt(30); section.top_margin = section.bottom_margin = Pt(36)
    paragraph = document.add_paragraph(); paragraph.paragraph_format.space_before = Pt(0); paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text); run.font.name = 'Noto Sans CJK ' + suffix; run.font.size = Pt(11)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK ' + suffix)
    temporary = path.with_suffix('.unpacked.docx'); document.save(temporary)
    with ZipFile(temporary) as source, ZipFile(path, 'w') as target:
        for name in source.namelist():
            target.writestr(ZipInfo(name, (2000, 1, 1, 0, 0, 0)), source.read(name), compress_type=ZIP_DEFLATED)
    temporary.unlink()


def main():
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args(); output = args.output.resolve()
    fixtures = (Path(__file__).resolve().parents[1] / 'tests/wordnet.Conversion.Tests/Fixtures').resolve()
    if output.exists() or output == fixtures or fixtures in output.parents:
        parser.error('Choose a new scratch directory outside the fixture tree.')
    if docx.__version__ != '1.2.0':
        raise RuntimeError('Requires python-docx 1.2.0.')
    output.mkdir(parents=True)
    for name, (suffix, text) in CASES.items():
        create(output / (name + '.docx'), suffix, text)
    print(json.dumps({path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(output.iterdir())}))


if __name__ == '__main__':
    main()
