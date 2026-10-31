"""Author synthetic rotated text-box fixtures with python-docx and Pillow."""
import argparse
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt
from PIL import Image


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def grid():
    image = Image.new('RGB', (8, 4))
    for x in range(8):
        for y in range(4):
            image.putpixel((x, y), (220, 40, 40) if x < 4 else (40, 70, 220))
    data = BytesIO(); image.save(data, format='PNG'); data.seek(0); return data


def create(path, angle, linked_image):
    document = docx.Document()
    document.core_properties.author = 'Synthetic fixture'
    document.core_properties.created = document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    section = document.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    document.add_paragraph('Before')
    owner = document.add_paragraph()
    drawing = parse_xml('''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
 <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
  <wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>1270000</wp:posOffset></wp:positionH>
  <wp:positionV relativeFrom="page"><wp:posOffset>1270000</wp:posOffset></wp:positionV><wp:extent cx="1778000" cy="1016000"/>
  <wp:wrapNone/><wp:docPr id="1" name="Synthetic rotated box"/><a:graphic>
   <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr>
    <a:xfrm rot="%d"><a:off x="0" y="0"/><a:ext cx="1778000" cy="1016000"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="00FFFF"/></a:solidFill><a:ln><a:noFill/></a:ln>
   </wps:spPr><wps:txbx><w:txbxContent/></wps:txbx><wps:bodyPr lIns="91440" rIns="91440" tIns="45720" bIns="45720"><a:noAutofit/></wps:bodyPr></wps:wsp></a:graphicData></a:graphic>
 </wp:anchor></w:drawing>''' % angle)
    content = drawing.find('.//' + qn('w:txbxContent'))
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0); paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if linked_image:
        rid = document.part.relate_to('https://example.test/callout', RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'), rid)
        run = OxmlElement('w:r'); properties = OxmlElement('w:rPr')
        fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'), 'Noto Sans'); fonts.set(qn('w:hAnsi'), 'Noto Sans')
        size = OxmlElement('w:sz'); size.set(qn('w:val'), '22'); properties.extend((fonts, size)); run.append(properties)
        text = OxmlElement('w:t'); text.text = 'Linked'; run.append(text); hyperlink.append(run); paragraph._p.append(hyperlink)
        run = paragraph.add_run(); run.add_picture(grid(), width=Pt(40), height=Pt(20))
        run._r.find('.//' + qn('wp:docPr')).set('id', '2')
    else:
        run = paragraph.add_run('Rotate'); run.font.name = 'Noto Sans'; run.font.size = Pt(11)
    content.append(paragraph._p); owner.add_run()._r.append(drawing)
    document.add_paragraph('After')
    temporary = path.with_suffix('.unpacked.docx'); document.save(temporary)
    with ZipFile(temporary) as source, ZipFile(path, 'w') as output:
        for name in source.namelist():
            output.writestr(ZipInfo(name, (2000, 1, 1, 0, 0, 0)), source.read(name), compress_type=ZIP_DEFLATED)
    temporary.unlink()


def main():
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument('--output', type=Path, required=True)
    output = parser.parse_args().output.resolve(); fixtures = (Path(__file__).resolve().parents[1] / 'tests/wordnet.Conversion.Tests/Fixtures').resolve()
    if output.exists() or output == fixtures or fixtures in output.parents:
        parser.error('Choose a new scratch directory outside the fixture tree.')
    if docx.__version__ != '1.2.0':
        raise RuntimeError('Requires python-docx 1.2.0.')
    output.mkdir(parents=True)
    for name, angle, linked in (('plain-30', 1800000, False), ('plain-45', 2700000, False),
                                ('plain-fractional', 2700001, False), ('linked-image-45', 2700000, True),
                                ('linked-image-90', 5400000, True)):
        create(output / (name + '.docx'), angle, linked)
    print(json.dumps({'producer': 'python-docx 1.2.0 and Pillow 12.3.0, synthetic DrawingML',
                      'generatorSha256': digest(Path(__file__)), 'documents': {p.name: digest(p) for p in sorted(output.iterdir())}}))


if __name__ == '__main__':
    main()
