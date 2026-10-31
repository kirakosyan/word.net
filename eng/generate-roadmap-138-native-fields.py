"""Create independent synthetic DOCX scalar-field inputs with python-docx."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def create(path, title, cache):
    document = docx.Document()
    document.core_properties.title = title
    document.core_properties.author = 'Synthetic fixture'
    document.core_properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    normal = document.styles['Normal']
    normal.font.name = 'Noto Sans'
    normal.font.size = Pt(11)
    paragraph = document.add_paragraph()
    paragraph.add_run('Label: ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'DOCPROPERTY Title')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = cache
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    paragraph.add_run(' End 42')
    temporary = path.with_suffix('.unpacked.docx')
    document.save(temporary)
    # Normalize only ZIP metadata; all OOXML parts are authored by python-docx.
    with ZipFile(temporary) as source, ZipFile(path, 'w') as output:
        for name in source.namelist():
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            output.writestr(info, source.read(name), compress_type=ZIP_DEFLATED)
    temporary.unlink()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True, type=Path)
    directory = parser.parse_args().output.resolve()
    fixtures = (Path(__file__).resolve().parents[1] / 'tests/wordnet.Conversion.Tests/Fixtures').resolve()
    if directory == fixtures or fixtures in directory.parents:
        parser.error('Use a scratch directory outside the fixture tree.')
    if docx.__version__ != '1.2.0':
        raise RuntimeError('This fixture requires python-docx 1.2.0.')
    directory.mkdir(parents=True, exist_ok=True)
    outputs = {}
    for name, title in [('hindi', 'क्षि'), ('arabic', 'مرحبا')]:
        for suffix, cache in [('stale', 'old-cache'), ('reference', title)]:
            output = directory / f'{name}-{suffix}.docx'
            if output.exists():
                raise FileExistsError('Choose a fresh output directory.')
            create(output, title, cache)
            outputs[output.name] = digest(output)
    manifest = {'producer': 'python-docx 1.2.0 (MIT), synthetic fields and core title',
                'generatorSha256': digest(Path(__file__)), 'documents': outputs}
    (directory / 'manifest.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + '\n', encoding='utf-8', newline='\n')
    print(json.dumps(manifest, ensure_ascii=False))


if __name__ == '__main__':
    main()
